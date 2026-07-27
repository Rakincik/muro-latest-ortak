import os
import sys
import re
import uuid
import datetime
import argparse
import openpyxl
import psycopg2

def normalize(name_str):
    if not name_str:
        return ""
    # Turkish char replacement
    char_map = {
        'ı': 'i', 'İ': 'i', 'I': 'i',
        'ş': 's', 'Ş': 's',
        'ğ': 'g', 'Ğ': 'g',
        'ü': 'u', 'Ü': 'u',
        'ö': 'o', 'Ö': 'o',
        'ç': 'c', 'Ç': 'c'
    }
    s = str(name_str).strip().lower()
    for tr, eng in char_map.items():
        s = s.replace(tr, eng)
    s = re.sub(r'[^a-z0-9\s]', '', s)
    return " ".join(s.split())

def parse_name(full_name):
    parts = str(full_name).strip().split()
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return " ".join(parts[:-1]), parts[-1]

def clean_phone_number(phone):
    if not phone:
        return None
    s = re.sub(r'\D', '', str(phone))
    # Strip leading 90 or 0
    if len(s) == 11 and s.startswith('0'):
        s = s[1:]
    elif len(s) == 12 and s.startswith('90'):
        s = s[2:]
    elif len(s) == 13 and s.startswith('+90'):
        s = s[3:]
    if len(s) == 10 and s.startswith('5'):
        return s
    return None

def generate_mock_phone(name):
    # Generates a unique 10-digit phone number starting with 111 based on name hash
    import hashlib
    h = hashlib.md5(normalize(name).encode('utf-8')).hexdigest()
    suffix = str(int(h, 16))[:7]
    # Ensure it has exactly 7 digits
    while len(suffix) < 7:
        suffix += "0"
    return "111" + suffix

def generate_mock_email(name):
    parts = normalize(name).split()
    if not parts:
        return f"ogrenci_{str(uuid.uuid4())[:8]}@3u.muro.click"
    if len(parts) == 1:
        return f"{parts[0]}@3u.muro.click"
    return f"{parts[0]}.{parts[-1]}@3u.muro.click"

def generate_password(name, phone):
    # adı + . + telefonun_son_2_hanesi + . + soyadının_baş_harfi
    # e.g., adile.85.y
    first, last = parse_name(name)
    first_clean = normalize(first).split()
    first_val = first_clean[0] if first_clean else "ogrenci"
    
    last_clean = normalize(last)
    last_char = last_clean[0] if last_clean else "x"
    
    phone_val = str(phone)
    phone_suffix = phone_val[-2:] if len(phone_val) >= 2 else "00"
    
    return f"{first_val}.{phone_suffix}.{last_char}"

def main():
    parser = argparse.ArgumentParser(description="3U LMS - 4T Groups & Students Import Tool")
    parser.add_argument('--host', default='localhost', help='PostgreSQL Host')
    parser.add_argument('--port', default='5432', help='PostgreSQL Port')
    parser.add_argument('--dbname', default='muro_dev', help='PostgreSQL Database Name')
    parser.add_argument('--user', default='muro_user', help='PostgreSQL User')
    parser.add_argument('--password', help='PostgreSQL Password')
    parser.add_argument('--execute', action='store_true', help='Commit changes to the database')
    parser.add_argument('--truncate', action='store_true', help='Clear existing Groups & Memberships before importing')
    
    args = parser.parse_args()
    
    conn_params = {
        'host': args.host,
        'port': args.port,
        'dbname': args.dbname,
        'user': args.user
    }
    if args.password:
        conn_params['password'] = args.password
    else:
        # Default local password
        conn_params['password'] = 'muro_pass_2024'

    script_dir = os.path.dirname(os.path.abspath(__file__))
    # Folder is inside "4T Ana Gruplar ve Öğrencileri"
    folder_path = os.path.abspath(os.path.join(script_dir, "..", "4T Ana Gruplar ve Öğrencileri"))
    
    print("============================================================")
    print("         3U LMS - 4T GROUPS & STUDENTS MIGRATION")
    print("============================================================")
    print(f"[*] Target Folder: {folder_path}")
    if not os.path.exists(folder_path):
        print(f"[!] Error: Directory '{folder_path}' does not exist.")
        sys.exit(1)
        
    files = [f for f in os.listdir(folder_path) if f.endswith('.xlsx')]
    print(f"[*] Found {len(files)} Excel files to process.")
    
    # Structure: group_key -> { 'name': str, 'parent_name': str or None, 'students': list }
    parsed_groups = {}
    all_students_global = {} # norm_name -> { name, phone, email, password, groups: set }
    
    for filename in sorted(files):
        file_path = os.path.join(folder_path, filename)
        group_name = filename[:-5].strip()
        group_key = group_name.lower()
        
        # Load workbook
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
        except Exception as e:
            print(f"[!] Warning: Failed to load '{filename}': {e}")
            continue
            
        student_list = []
        # Row 1 is title, Row 2 is headers, Row 3 onwards is data
        if len(rows) > 2:
            for r in rows[2:]:
                name = r[1]
                phone = r[2]
                email = r[3]
                
                if name:
                    name_str = str(name).strip()
                    norm_name = normalize(name_str)
                    
                    cleaned_phone = clean_phone_number(phone)
                    if not cleaned_phone:
                        cleaned_phone = generate_mock_phone(name_str)
                        
                    cleaned_email = str(email).strip().lower() if email else None
                    if not cleaned_email or "@" not in cleaned_email:
                        cleaned_email = generate_mock_email(name_str)
                        
                    pwd = generate_password(name_str, cleaned_phone)
                    
                    student_list.append(norm_name)
                    
                    if norm_name not in all_students_global:
                        all_students_global[norm_name] = {
                            'name': name_str,
                            'phone': cleaned_phone,
                            'email': cleaned_email,
                            'password': pwd,
                            'groups': set()
                        }
                    all_students_global[norm_name]['groups'].add(group_key)
                    
        parsed_groups[group_key] = {
            'raw_key': group_name,
            'name': group_name,
            'parent_name': None,
            'students': student_list
        }
        
    print(f"[+] Loaded successfully:")
    print(f"    - Unique groups parsed    : {len(parsed_groups)}")
    print(f"    - Total unique students   : {len(all_students_global)}")

    # Connect to Database
    print(f"\n[*] Connecting to PostgreSQL database...")
    try:
        conn = psycopg2.connect(**conn_params)
        cur = conn.cursor()
        print("[+] Connected successfully!")
    except Exception as e:
        print(f"[!] Database Connection Failed: {e}")
        sys.exit(1)

    try:
        # Check if TenantMemberships table exists in this database
        cur.execute("""
            SELECT EXISTS (
                SELECT FROM information_schema.tables 
                WHERE table_schema = 'public' 
                AND table_name = 'TenantMemberships'
            );
        """)
        has_tenant_memberships = cur.fetchone()[0]

        # Cache existing CourseGroups mappings BEFORE any changes
        print("[*] Retrieving existing database records...")
        cur.execute("""
            SELECT g."Name", c."Id", cg."Mode"
            FROM "CourseGroups" cg
            JOIN "Groups" g ON cg."GroupId" = g."Id"
            JOIN "Courses" c ON cg."CourseId" = c."Id";
        """)
        raw_cg_mappings = cur.fetchall()
        existing_cg_mappings = {} # group_name_lower -> list of (course_uuid, mode)
        for g_name, c_uuid, cg_mode in raw_cg_mappings:
            g_name_lower = g_name.strip().lower()
            if g_name_lower not in existing_cg_mappings:
                existing_cg_mappings[g_name_lower] = []
            existing_cg_mappings[g_name_lower].append((c_uuid, cg_mode))
        print(f"[+] Cached {len(raw_cg_mappings)} existing course-group mappings from database to preserve them.")

        if args.truncate:
            print("\n[!] WARNING: Truncate requested. Emptying Groups and Memberships...")
            if not args.execute:
                print("    - Dry-run simulation: Truncate queries would be executed.")
            else:
                cur.execute('TRUNCATE TABLE "GroupMembers" CASCADE;')
                cur.execute('TRUNCATE TABLE "CourseGroups" CASCADE;')
                cur.execute('TRUNCATE TABLE "Groups" CASCADE;')
                print("    - Database Groups, CourseGroups, and GroupMembers truncated successfully!")

        # Retrieve existing records after potential truncate
        cur.execute('SELECT "Id", "Name", "Description" FROM "Groups" WHERE "IsDeleted" = False;')
        db_groups = cur.fetchall()
        db_groups_by_name = {r[1].lower().strip(): r[0] for r in db_groups}
        
        cur.execute('SELECT "Id", "Email", "Phone", "Username" FROM "Users";')
        db_users = cur.fetchall()
        db_users_by_email = {r[1].lower().strip(): r[0] for r in db_users if r[1]}
        db_users_by_phone = {r[2].strip(): r[0] for r in db_users if r[2]}
        db_users_by_username = {r[3].lower().strip(): r[0] for r in db_users if r[3]}

        db_tenant_members = set()
        if has_tenant_memberships:
            cur.execute('SELECT "UserId" FROM "TenantMemberships";')
            db_tenant_members = set(r[0] for r in cur.fetchall())

        cur.execute('SELECT "UserId", "GroupId" FROM "GroupMembers";')
        db_group_members = set((r[0], r[1]) for r in cur.fetchall())

        print(f"[+] Loaded existing: {len(db_groups)} groups, {len(db_users)} users (TenantMemberships exists: {has_tenant_memberships}).")

        # Proposed changes summary
        groups_to_create = list(parsed_groups.keys())
        students_to_create = list(all_students_global.keys())
        
        print("\n" + "=" * 50)
        print("                  PROPOSED WORKPLAN")
        print("=" * 50)
        if args.truncate:
            print(" * CLEAR ALL EXISTING GROUPS AND ASSIGNMENTS (TRUNCATE)")
        print(f" 1. Create Groups     : {len(groups_to_create)} groups to verify/create")
        print(f" 2. Create Students   : {len(students_to_create)} total student accounts to evaluate")
        print(f" 3. Group Memberships : Assigning students to their respective groups")
        print("=" * 50)

        if not args.execute:
            print("\n[*] RUNNING IN DRY-RUN (SIMULATION) MODE. No database writes will be executed.")
            print("[*] To perform the actual import, run with: python import_3u.py --execute")
            if args.truncate:
                print("[*] Add --truncate if you also want to clean groups before import.")
            cur.close()
            conn.close()
            sys.exit(0)

        print("\n[*] Starting transaction...")

        # 1. Create Groups
        print("[1/4] Writing Groups...")
        group_uuid_map = {} # group_name_lower -> uuid
        groups_created = 0
        
        for g_key, g_info in parsed_groups.items():
            g_name = g_info['name']
            
            # Check if group already exists in database (matching by name only, case-insensitive!)
            cur.execute('SELECT "Id" FROM "Groups" WHERE LOWER("Name") = %s AND "IsDeleted" = False ORDER BY "CreatedAt" DESC;', (g_name.lower().strip(),))
            row = cur.fetchone()
            if row:
                db_uuid = row[0]
            else:
                db_uuid = str(uuid.uuid4())
                desc = f"Imported Group: {g_name}"
                cur.execute(
                    'INSERT INTO "Groups" ("Id", "Name", "Description", "ParentId", "IsDeleted", "CreatedAt") VALUES (%s, %s, %s, %s, %s, %s)',
                    (db_uuid, g_name, desc, None, False, datetime.datetime.utcnow())
                )
                groups_created += 1
            group_uuid_map[g_key] = db_uuid
            
        print(f"      -> Created: {groups_created} new groups, Reused: {len(parsed_groups) - groups_created} existing.")

        # 2. Re-link Courses to Groups (if they existed before truncate or if we have mappings in memory)
        print("[2/4] Restoring Course-Group mappings...")
        cg_restored = 0
        
        for g_key, g_uuid in group_uuid_map.items():
            # Check if this group name had any course mappings cached in memory
            # g_key is already lowercase name
            if g_key in existing_cg_mappings:
                for c_uuid, cg_mode in existing_cg_mappings[g_key]:
                    # Check if already linked (e.g. if we didn't truncate)
                    cur.execute('SELECT 1 FROM "CourseGroups" WHERE "CourseId" = %s AND "GroupId" = %s;', (c_uuid, g_uuid))
                    if not cur.fetchone():
                        cur.execute(
                            'INSERT INTO "CourseGroups" ("Id", "CourseId", "GroupId", "Mode", "AssignedAt") VALUES (%s, %s, %s, %s, %s)',
                            (str(uuid.uuid4()), c_uuid, g_uuid, cg_mode, datetime.datetime.utcnow())
                        )
                        cg_restored += 1
        print(f"      -> Restored: {cg_restored} course-group mappings.")

        # 3. Create Students
        print("[3/4] Creating Student accounts...")
        user_map = {} # norm_name -> uuid
        students_created = 0
        students_skipped = 0
        tm_created = 0

        for norm_name, stud in all_students_global.items():
            name = stud['name']
            phone = stud['phone']
            email = stud['email']
            pwd = stud['password']
            first, last = parse_name(name)

            phone_val = phone
            email_val = email

            u_id = None
            if email_val.lower().strip() in db_users_by_email:
                u_id = db_users_by_email[email_val.lower().strip()]
            elif phone_val in db_users_by_phone:
                u_id = db_users_by_phone[phone_val]
            elif phone_val in db_users_by_username:
                u_id = db_users_by_username[phone_val]

            if u_id:
                user_map[norm_name] = u_id
                students_skipped += 1
            else:
                u_id = str(uuid.uuid4())
                user_map[norm_name] = u_id
                cur.execute(
                    'INSERT INTO "Users" ("Id", "FirstName", "LastName", "Email", "Username", "Phone", "PasswordHash", "Role", "StudentType", "IsActive", "IsDeleted", "CreatedAt") VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)',
                    (u_id, first, last, email_val, phone_val, phone_val, pwd, 'Student', 'Active', True, False, datetime.datetime.utcnow())
                )
                db_users_by_email[email_val.lower().strip()] = u_id
                db_users_by_phone[phone_val] = u_id
                db_users_by_username[phone_val] = u_id
                students_created += 1

            if has_tenant_memberships and u_id not in db_tenant_members:
                cur.execute(
                    'INSERT INTO "TenantMemberships" ("Id", "UserId", "Role", "Status", "JoinedAt") VALUES (%s, %s, %s, %s, %s)',
                    (str(uuid.uuid4()), u_id, 'Student', 'active', datetime.datetime.utcnow())
                )
                db_tenant_members.add(u_id)
                tm_created += 1
                
        print(f"      -> Created: {students_created} students, Associated Tenant: {tm_created}, Skipped: {students_skipped}.")

        # 4. Assign to Groups
        print("[4/4] Mapping Students to Groups...")
        gm_created = 0
        gm_skipped = 0
        
        for norm_name, stud in all_students_global.items():
            u_uuid = user_map.get(norm_name)
            if not u_uuid:
                continue
            for g_key in stud['groups']:
                g_uuid = group_uuid_map.get(g_key)
                if g_uuid:
                    if (u_uuid, g_uuid) not in db_group_members:
                        cur.execute(
                            'INSERT INTO "GroupMembers" ("Id", "UserId", "GroupId", "Role", "Status", "AddedAt") VALUES (%s, %s, %s, %s, %s, %s)',
                            (str(uuid.uuid4()), u_uuid, g_uuid, 2, 'active', datetime.datetime.utcnow())
                        )
                        db_group_members.add((u_uuid, g_uuid))
                        gm_created += 1
                    else:
                        gm_skipped += 1
        print(f"      -> Created: {gm_created} memberships, Skipped: {gm_skipped}.")

        # Commit Transaction
        print("\n[*] Committing database transaction...")
        conn.commit()
        print("[+] SUCCESS: Data migration completed and committed successfully!")

    except Exception as db_err:
        print(f"\n[!] Database Transaction Error: {db_err}")
        print("[*] Rolling back changes to ensure data integrity...")
        conn.rollback()
        print("[*] Rollback completed. No changes saved.")
        sys.exit(1)
    finally:
        cur.close()
        conn.close()

if __name__ == '__main__':
    main()
